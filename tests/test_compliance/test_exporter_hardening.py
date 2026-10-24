"""Regression tests for the compliance exporter hardening wave (R30-R40).

Covers tenant isolation (R30/R31), risk-field population (R32, with a
100,000 loss day), spreadsheet formula injection (R33), export audit records
(R34), the half-open date window (R35), the exchange timezone basis (R36),
filename sanitization/uniqueness/atomicity (R37/R38), format whitelisting
(R39) and skipped-row accounting / up-front window validation (R40).
"""

from __future__ import annotations

import builtins
import csv
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest

import quant_platform.compliance.exporter as exporter_module
from quant_platform.compliance.exporter import ComplianceExporter
from quant_platform.core.context import TenantContext
from quant_platform.core.store import Store


@pytest.fixture
def store(tmp_path):
    return Store(str(tmp_path / "test.db"))


@pytest.fixture
def exporter(tmp_path, store):
    return ComplianceExporter(store, output_dir=str(tmp_path / "output"))


def _seed_trades(store, tenant_id=None, n=2, prefix="t"):
    for i in range(n):
        trade = {
            "trade_id": f"{prefix}{i}", "order_id": f"o{prefix}{i}",
            "code": "600519", "side": "buy", "quantity": 100,
            "price": 10.0, "executed_at": "2024-06-05T10:00:00",
        }
        if tenant_id is not None:
            trade["tenant_id"] = tenant_id
        store.save_trade(trade)


def _seed_loss_days(store):
    for timestamp, equity, market_value in (
        ("2024-06-03T15:00:00", 10_000_000.0, 5_000_000.0),
        ("2024-06-04T15:00:00", 9_900_000.0, 4_900_000.0),
        ("2024-06-05T15:00:00", 9_950_000.0, 4_950_000.0),
    ):
        store.save_pnl_snapshot({
            "timestamp": timestamp,
            "total_equity": equity,
            "cash": 5_000_000.0,
            "market_value": market_value,
        })


# ── R30/R31: tenant isolation ───────────────────────────────────────────


class TestTenantIsolation:
    def test_default_export_does_not_merge_tenants(self, store, exporter):
        _seed_trades(store, tenant_id="fund_a", n=2, prefix="a")
        _seed_trades(store, tenant_id="fund_b", n=3, prefix="b")
        _seed_trades(store, tenant_id=None, n=1, prefix="d")

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )

        assert result.tenant_id == "default"
        assert result.n_records == 1
        with open(result.file_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        assert [row[0] for row in rows[1:]] == ["d0"]

    def test_tenant_context_scopes_the_default_export(self, store, exporter):
        _seed_trades(store, tenant_id="fund_a", n=2, prefix="a")
        _seed_trades(store, tenant_id="fund_b", n=3, prefix="b")

        TenantContext.set_current(TenantContext(tenant_id="fund_a"))
        try:
            result = exporter.export_trade_log(
                start="2024-06-01", end="2024-06-30", format="csv",
            )
        finally:
            TenantContext.clear()

        assert result.tenant_id == "fund_a"
        assert result.n_records == 2

    def test_cross_tenant_export_requires_an_explicit_flag(
        self, store, exporter,
    ):
        _seed_trades(store, tenant_id="fund_a", n=2, prefix="a")
        _seed_trades(store, tenant_id="fund_b", n=3, prefix="b")

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
            include_all_tenants=True,
        )

        assert result.tenant_id == ""
        assert result.n_records == 5

    def test_tenant_and_all_tenants_together_are_rejected(
        self, store, exporter,
    ):
        _seed_trades(store, tenant_id="fund_a", n=1, prefix="a")

        with pytest.raises(ValueError, match="mutually exclusive"):
            exporter.export_trade_log(
                start="2024-06-01", end="2024-06-30", format="csv",
                tenant_id="fund_a", include_all_tenants=True,
            )

    def test_order_export_is_tenant_scoped_by_context(self, store, exporter):
        for tenant in ("fund_a", "fund_b"):
            store.save_order({
                "order_id": f"o-{tenant}", "tenant_id": tenant,
                "code": "600519", "side": "buy", "quantity": 100,
                "price": 10.0, "status": "filled",
                "created_at": "2024-06-05T09:30:00",
            })

        TenantContext.set_current(TenantContext(tenant_id="fund_b"))
        try:
            result = exporter.export_order_log(
                start="2024-06-01", end="2024-06-30", format="csv",
            )
        finally:
            TenantContext.clear()

        assert result.tenant_id == "fund_b"
        assert result.n_records == 1


class TestRiskLogTenantSchemaLimit:
    """pnl_history has no tenant_id column: the exporter must refuse, not
    stamp a shared file with one fund's label."""

    def test_explicit_tenant_is_refused(self, store, exporter):
        store.save_pnl_snapshot({
            "timestamp": "2024-06-03T15:00:00",
            "total_equity": 10_000_000.0, "cash": 5_000_000.0,
            "market_value": 5_000_000.0,
        })

        with pytest.raises(ValueError, match="no tenant_id"):
            exporter.export_risk_log(
                start="2024-06-01", end="2024-06-30", tenant_id="fund_a",
            )

    def test_whole_book_flag_is_required(self, store, exporter):
        store.save_pnl_snapshot({
            "timestamp": "2024-06-03T15:00:00",
            "total_equity": 10_000_000.0, "cash": 5_000_000.0,
            "market_value": 5_000_000.0,
        })

        # No context, no explicit tenant: the file is whole-book and must be
        # acknowledged as such — it is never stamped as tenant-scoped data.
        with pytest.raises(ValueError, match="include_all_tenants"):
            exporter.export_risk_log(start="2024-06-01", end="2024-06-30")

        # An explicit flag wins even under a non-default context, and the
        # result says "all", not the context tenant.
        TenantContext.set_current(TenantContext(tenant_id="fund_a"))
        try:
            result = exporter.export_risk_log(
                start="2024-06-01", end="2024-06-30",
                include_all_tenants=True,
            )
        finally:
            TenantContext.clear()

        assert result.n_records == 1
        assert result.tenant_id == ""


# ── R32: risk-field population ──────────────────────────────────────────


class TestRiskFieldPopulation:
    def test_pnl_writer_derives_the_five_risk_columns(self, store):
        _seed_loss_days(store)

        rows = store.get_pnl_history(
            start_date="2024-06-01", end_date="2024-06-30",
        )
        assert len(rows) == 3

        loss = rows[1]
        assert loss["daily_pnl"] == pytest.approx(-100_000.0)
        assert loss["daily_pnl_pct"] == pytest.approx(-0.01)
        assert loss["cumulative_pnl"] == pytest.approx(-100_000.0)
        assert loss["max_drawdown"] == pytest.approx(-0.01)
        assert loss["sharpe_ratio"] == 0.0  # only one return so far

        recovery = rows[2]
        expected_returns = [-0.01, 50_000.0 / 9_900_000.0]
        mean = sum(expected_returns) / len(expected_returns)
        variance = sum((r - mean) ** 2 for r in expected_returns) / (
            len(expected_returns) - 1
        )
        expected_sharpe = mean / (variance ** 0.5) * (252 ** 0.5)

        assert recovery["daily_pnl"] == pytest.approx(50_000.0)
        assert recovery["daily_pnl_pct"] == pytest.approx(50_000.0 / 9_900_000.0)
        assert recovery["cumulative_pnl"] == pytest.approx(-50_000.0)
        assert recovery["max_drawdown"] == pytest.approx(-0.01)
        assert recovery["sharpe_ratio"] == pytest.approx(expected_sharpe)

    def test_explicit_caller_values_still_win(self, store):
        store.save_pnl_snapshot({
            "timestamp": "2024-06-03T15:00:00",
            "total_equity": 10_000_000.0, "cash": 5_000_000.0,
            "market_value": 5_000_000.0,
        })
        store.save_pnl_snapshot({
            "timestamp": "2024-06-04T15:00:00",
            "total_equity": 9_900_000.0, "cash": 5_000_000.0,
            "market_value": 4_900_000.0,
            "daily_pnl": 123.0, "daily_pnl_pct": 0.5,
            "cumulative_pnl": 7.0, "max_drawdown": -0.5,
            "sharpe_ratio": 2.0,
        })

        rows = store.get_pnl_history(
            start_date="2024-06-01", end_date="2024-06-30",
        )
        row = rows[1]
        assert row["daily_pnl"] == 123.0
        assert row["daily_pnl_pct"] == 0.5
        assert row["cumulative_pnl"] == 7.0
        assert row["max_drawdown"] == -0.5
        assert row["sharpe_ratio"] == 2.0

    def test_risk_log_exports_the_100k_loss_day(self, store, exporter):
        _seed_loss_days(store)

        result = exporter.export_risk_log(
            start="2024-06-01", end="2024-06-30", format="csv",
            include_all_tenants=True,
        )

        with open(result.file_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))

        assert len(rows) == 4  # header + 3 snapshots
        loss_row = rows[2]  # 2024-06-04
        assert loss_row[0] == "2024-06-04T15:00:00"
        assert float(loss_row[4]) == pytest.approx(-100_000.0)
        assert float(loss_row[5]) == pytest.approx(-0.01)
        assert float(loss_row[6]) == pytest.approx(-100_000.0)
        assert float(loss_row[8]) == pytest.approx(-0.01)


# ── R33: spreadsheet formula injection ──────────────────────────────────


class TestFormulaInjection:
    def test_csv_cells_are_neutralized(self, store, exporter):
        store.save_trade({
            "trade_id": "=1+1", "order_id": "+SUM(A1:A2)",
            "tenant_id": "@cmd", "code": "600519", "side": "-1",
            "quantity": 100, "price": 10.0,
            "executed_at": "2024-06-05T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
            tenant_id="@cmd",
        )

        with open(result.file_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        row = rows[1]
        assert row[0] == "'=1+1"
        assert row[1] == "'+SUM(A1:A2)"
        assert row[2] == "'@cmd"      # tenant column is data too
        assert row[4] == "'-1"
        assert row[5] == "100"  # numeric cells stay numeric

    def test_csv_leading_control_character_is_neutralized(
        self, store, exporter,
    ):
        store.save_trade({
            "trade_id": "t1", "order_id": "o1", "code": "600519",
            "side": "\t=1+1", "quantity": 100, "price": 10.0,
            "executed_at": "2024-06-05T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )

        with open(result.file_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        assert rows[1][4] == "'\t=1+1"

    def test_xlsx_cells_are_stored_as_text(self, store, exporter):
        from openpyxl import load_workbook

        store.save_trade({
            "trade_id": "=1+1", "order_id": "o1", "code": "600519",
            "side": "@cmd", "quantity": 100, "price": 10.0,
            "executed_at": "2024-06-05T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="xlsx",
        )

        ws = load_workbook(result.file_path).active
        assert ws["A2"].value == "'=1+1"
        assert ws["A2"].data_type == "s"       # text, not a formula ('f')
        assert ws["E2"].value == "'@cmd"
        assert ws["E2"].data_type == "s"
        assert ws["F2"].value == 100
        assert ws["F2"].data_type == "n"       # numbers remain numeric

    def test_xlsx_illegal_control_characters_do_not_block_the_export(
        self, store, exporter,
    ):
        """openpyxl rejects control characters anywhere in a cell; dirty
        data must not permanently block the compliance export."""
        from openpyxl import load_workbook

        store.save_trade({
            "trade_id": "t1", "order_id": "o1", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "commission": "\x01=1+1",
            "executed_at": "2024-06-05T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="xlsx",
        )

        ws = load_workbook(result.file_path).active
        assert ws["H2"].value == "'=1+1"   # control char removed, '=' escaped
        assert ws["H2"].data_type == "s"


# ── R34: export audit trail ─────────────────────────────────────────────


class TestExportAudit:
    def test_success_is_audited(self, store, exporter):
        _seed_trades(store, n=2)

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )

        events = store.get_events(topic="audit.export", limit=10)
        assert len(events) == 1
        data = events[0]["data"]
        assert data["action"] == "export"
        assert data["result"] == "success"
        assert data["component"] == "compliance.exporter"
        details = data["details"]
        assert details["export"] == "trade_log"
        assert details["tenant_id"] == "default"
        assert details["n_records"] == 2
        assert details["n_skipped"] == 0
        assert details["start"] == "2024-06-01"
        assert details["end"] == "2024-06-30"
        assert details["file_path"] == result.file_path

    def test_failure_is_audited_and_reraised(self, store, tmp_path):
        class ExplodingQueryStore:
            def __init__(self, inner):
                self.inner = inner

            def get_trades_for_export(self, **kwargs):
                raise RuntimeError("db read failed")

            def log_event(self, event):
                self.inner.log_event(event)

        exporter = ComplianceExporter(
            ExplodingQueryStore(store), output_dir=str(tmp_path / "out"),
        )

        with pytest.raises(RuntimeError, match="db read failed"):
            exporter.export_trade_log(
                start="2024-06-01", end="2024-06-30", format="csv",
            )

        events = store.get_events(topic="audit.export", limit=10)
        assert len(events) == 1
        data = events[0]["data"]
        assert data["result"] == "failure"
        assert "RuntimeError: db read failed" in data["details"]["error"]
        assert data["details"]["file_path"] == ""

    def test_refused_export_is_audited(self, store, exporter):
        """A rejected format/date is still a security-relevant attempt."""
        with pytest.raises(ValueError, match="Unsupported export format"):
            exporter.export_trade_log(
                start="2024-06-01", end="2024-06-30", format="parquet",
            )

        events = store.get_events(topic="audit.export", limit=10)
        assert len(events) == 1
        data = events[0]["data"]
        assert data["result"] == "failure"
        # The scope was never resolved, so the record says "unknown" rather
        # than implying a whole-book export was attempted.
        assert data["details"]["tenant_id"] == "unknown"
        assert "Unsupported export format" in data["details"]["error"]

    def test_audit_row_is_scoped_to_the_exported_tenant(self, store, exporter):
        """A tenant-scoped audit query must find exports of that tenant's
        data even when the acting context is different."""
        _seed_trades(store, tenant_id="fund_a", n=1, prefix="a")

        exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
            tenant_id="fund_a",
        )

        events = store.get_events(
            topic="audit.export", tenant_id="fund_a", limit=10,
        )
        assert len(events) == 1
        assert events[0]["data"]["details"]["tenant_id"] == "fund_a"
        # The acting tenant is preserved separately for attribution.
        assert events[0]["data"]["actor_tenant"] == "default"


# ── R35/R36: date window and timezone basis ─────────────────────────────


class TestDateWindow:
    def test_final_fractional_second_is_included(self, store, exporter):
        store.save_trade({
            "trade_id": "frac", "order_id": "o1", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "executed_at": "2024-03-31T23:59:59.500000",
        })

        result = exporter.export_trade_log(
            start="2024-03-01", end="2024-03-31", format="csv",
        )

        assert result.n_records == 1

    def test_risk_log_includes_fractional_end_of_day(self, store, exporter):
        store.save_pnl_snapshot({
            "timestamp": "2024-03-31T23:59:59.999999",
            "total_equity": 10_000_000.0, "cash": 5_000_000.0,
            "market_value": 5_000_000.0,
        })

        result = exporter.export_risk_log(
            start="2024-03-01", end="2024-03-31", format="csv",
            include_all_tenants=True,
        )

        assert result.n_records == 1

    def test_next_midnight_is_excluded(self, store, exporter):
        store.save_trade({
            "trade_id": "next", "order_id": "o1", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "executed_at": "2024-04-01T00:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-03-01", end="2024-03-31", format="csv",
        )

        assert result.n_records == 0

    def test_aware_timestamps_use_exchange_tz_not_host(
        self, store, exporter, monkeypatch,
    ):
        # 17:00 UTC on Mar 31 is 01:00 on Apr 1 in Shanghai.
        store.save_trade({
            "trade_id": "tz", "order_id": "o1", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "executed_at": "2024-03-31T17:00:00+00:00",
        })

        monkeypatch.setattr(exporter_module, "EXCHANGE_TZ", ZoneInfo("UTC"))
        utc_result = exporter.export_trade_log(
            start="2024-03-01", end="2024-03-31", format="csv",
        )
        assert utc_result.n_records == 1

        monkeypatch.setattr(
            exporter_module, "EXCHANGE_TZ", ZoneInfo("Asia/Shanghai"),
        )
        shanghai_result = exporter.export_trade_log(
            start="2024-03-01", end="2024-03-31", format="csv",
        )
        assert shanghai_result.n_records == 0


# ── R37/R38: filenames ──────────────────────────────────────────────────


class TestFilenames:
    def test_tenant_is_sanitized_for_the_filename(self, store, exporter):
        tenant = "fund/a b"
        store.save_trade({
            "trade_id": "t1", "order_id": "o1", "tenant_id": tenant,
            "code": "600519", "side": "buy", "quantity": 100,
            "price": 10.0, "executed_at": "2024-06-05T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
            tenant_id=tenant,
        )

        assert result.n_records == 1
        assert Path(result.file_path).exists()
        assert "fund_a_b" in Path(result.file_path).name

    def test_repeated_exports_do_not_overwrite(self, store, exporter):
        _seed_trades(store, n=1)

        first = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )
        second = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )

        assert first.file_path != second.file_path
        assert Path(first.file_path).exists()
        assert Path(second.file_path).exists()
        assert not list(
            Path(first.file_path).parent.glob("*.tmp-*")
        )


# ── R39: format whitelist and fallback reporting ────────────────────────


class TestFormatHandling:
    def test_unknown_format_is_refused(self, exporter):
        with pytest.raises(ValueError, match="Unsupported export format"):
            exporter.export_trade_log(
                start="2024-06-01", end="2024-06-30", format="parquet",
            )

    def test_xlsx_fallback_reports_the_actual_format(
        self, store, exporter, monkeypatch,
    ):
        _seed_trades(store, n=1)
        original_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("simulated missing openpyxl")
            return original_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="xlsx",
        )

        assert result.format == "csv"
        assert result.file_path.endswith(".csv")


# ── R40: up-front validation and skipped-row accounting ─────────────────


class TestWindowValidationAndSkippedRows:
    def test_malformed_dates_raise_before_any_query(self, exporter):
        with pytest.raises(ValueError, match="start must be a YYYY-MM-DD"):
            exporter.export_trade_log(
                start="06/01/2024", end="2024-06-30",
            )
        with pytest.raises(ValueError, match="end must be a YYYY-MM-DD"):
            exporter.export_trade_log(
                start="2024-06-01", end="2024-06-31",
            )
        with pytest.raises(ValueError, match="after"):
            exporter.export_trade_log(
                start="2024-07-01", end="2024-06-30",
            )

    def test_extreme_dates_raise_value_error_not_overflow(self, exporter):
        with pytest.raises(ValueError, match="supported export range"):
            exporter.export_trade_log(
                start="0001-01-01", end="2024-06-30",
            )
        with pytest.raises(ValueError, match="supported export range"):
            exporter.export_risk_log(
                start="2024-01-01", end="9999-12-31",
                include_all_tenants=True,
            )

    def test_skipped_records_are_counted(self, store, exporter):
        store.save_trade({
            "trade_id": "bad", "order_id": "o1", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "executed_at": "2024-06-15Tnot-a-time",
        })
        store.save_trade({
            "trade_id": "ok", "order_id": "o2", "code": "600519",
            "side": "buy", "quantity": 100, "price": 10.0,
            "executed_at": "2024-06-15T10:00:00",
        })

        result = exporter.export_trade_log(
            start="2024-06-01", end="2024-06-30", format="csv",
        )

        assert result.n_records == 1
        assert result.n_skipped == 1

    def test_filter_counts_missing_and_unparseable_fields(self, exporter):
        filtered, skipped = exporter._filter_by_date(
            [
                {"ts": ""},
                {"ts": None},
                {"ts": "nope"},
                {"ts": "2024-06-15T00:00:00"},
            ],
            "ts", "2024-06-01", "2024-06-30",
        )

        assert len(filtered) == 1
        assert skipped == 3


# ── R42 companion: dates must be validated for the risk log too ─────────


def test_risk_log_rejects_malformed_window(exporter):
    with pytest.raises(ValueError, match="YYYY-MM-DD"):
        exporter.export_risk_log(start="2024-6-1", end="2024-06-30")
